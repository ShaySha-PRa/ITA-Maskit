"""生成「敏感信息不在规整列」的 Excel/CSV 测试数据集。

覆盖当前表格引擎：列映射校验、整格值级检测、格内强特征扫描、人员清单。

用法（项目根目录）:
  python scripts/gen_messy_testdata.py
  # 产出 testdata/messy/
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "testdata" / "messy"

NAMES = ["张伟", "李娜", "欧阳修", "司马光", "东方不败", "王芳", "陈静"]
EXTRA_NAMES = ["赵磊", "黄敏", "周杰"]
EMAILS = [
    "alice@corp.example",
    "bob.liu@company.cn",
    "carol@internal.local",
    "david.zhang@outlook.com",
    "eva.li@corp.example",
]
IPS = ["10.1.2.3", "192.168.0.88", "172.16.9.41", "10.20.30.40"]
ID_CARDS = [
    "110101199003077777",
    "310101198801011234",
    "44030119951212567X",
    "110101199001011234",
]
BANK_CARDS = [
    "6222021234567890123",
    "6228480012345678901",
    "621700123456789012",
    "6222029876543210987",
]
PHONES = ["13800001111", "139-0000-2222", "18612345678"]  # 格内默认不扫


def _style_header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.font = Font(bold=True)


def _widen(ws, width: int = 18) -> None:
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = width


def write_people_csv(path: Path) -> None:
    path.write_text("name\n" + "\n".join(NAMES) + "\n", encoding="utf-8")


def sheet_messy_notes(wb: Workbook) -> None:
    ws = wb.create_sheet("混乱备注", 0)
    ws.append(["序号", "事项", "备注", "说明", "附件摘要", "经办"])
    _style_header(ws)
    for r in [
        [1, "服务器巡检", EMAILS[0], "正常", IPS[0], "运维一组"],
        [2, "报销单", "见说明", ID_CARDS[0], "发票已传", NAMES[0]],
        [3, "供应商对账", BANK_CARDS[0], EMAILS[1], "本期结清", NAMES[1]],
        [4, "VPN 开通", IPS[1], "临时账号", EMAILS[2], NAMES[2]],
        [5, "员工入职", NAMES[3], ID_CARDS[1], BANK_CARDS[1], "人事"],
        [6, "日志排查", "普通文字无敏感", "流程OK", "无", EXTRA_NAMES[0]],
        [7, "跨境汇款", BANK_CARDS[2], NAMES[4], EMAILS[3], "财务"],
        [8, "会议室预约", "策划部", "合计", "其他", "一队"],
    ]:
        ws.append(r)
    _widen(ws)


def sheet_embedded_prose(wb: Workbook) -> None:
    ws = wb.create_sheet("嵌套正文")
    ws.append(["日期", "记录类型", "正文"])
    _style_header(ws)
    prose = [
        f"申请人：{NAMES[0]}，联系邮箱 {EMAILS[0]}，请审批。",
        f"故障主机 IP 为 {IPS[0]}，值班 {NAMES[1]}。",
        f"银行卡号 {BANK_CARDS[0]} 已登记，持卡人{NAMES[2]}。",
        f"身份证 {ID_CARDS[0]} 复印件已归档（经办 {NAMES[3]}）。",
        "本周无敏感事项，仅例会纪要。",
        f"供应商：亚玛芬体育，对接人 {NAMES[5]}，邮箱{EMAILS[1]}。",
    ]
    for i, text in enumerate(prose, start=1):
        ws.append([f"2026-03-{i:02d}", "纪要" if i % 2 else "工单", text])
    ws.column_dimensions["C"].width = 60
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True)


def sheet_wrong_headers(wb: Workbook) -> None:
    ws = wb.create_sheet("无语义列名")
    ws.append(["字段1", "字段2", "字段3", "字段4", "字段5"])
    _style_header(ws)
    ws.append([NAMES[0], EMAILS[0], ID_CARDS[0], BANK_CARDS[0], IPS[0]])
    ws.append([NAMES[1], EMAILS[1], ID_CARDS[1], BANK_CARDS[1], IPS[1]])
    ws.append([EXTRA_NAMES[1], "普通文本", "2.5", "2024.1.1", "=SUM(1,2)"])
    ws.append([NAMES[6], EMAILS[2], ID_CARDS[2], BANK_CARDS[2], IPS[2]])


def sheet_title_then_table(wb: Workbook) -> None:
    ws = wb.create_sheet("标题错位")
    ws.merge_cells("A1:E1")
    ws["A1"] = "2026 Q1 敏感信息抽查底稿（内部）"
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["编号", "内容A", "内容B", "内容C", "内容D"])
    ws.append(["R1", EMAILS[3], NAMES[0], ID_CARDS[0], "备注无敏感"])
    ws.append(["R2", IPS[2], NAMES[2], BANK_CARDS[0], EXTRA_NAMES[2]])


def sheet_athlete_payroll(wb: Workbook) -> None:
    ws = wb.create_sheet("选手收入")
    ws.append(["序号", "俱乐部", "选手", "教练", "应付", "备注邮箱"])
    _style_header(ws)
    ws.append([1, "一队", NAMES[0], NAMES[1], "12000", EMAILS[0]])
    ws.append([2, "二队", NAMES[2], NAMES[3], "8000", EMAILS[1]])
    ws.append([3, "青年队", NAMES[4], EXTRA_NAMES[0], "5000", "无"])
    ws.append([4, "合计", "合计", "其他", "25000", "策划部"])


def build_workbook_mixed() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    sheet_messy_notes(wb)
    sheet_embedded_prose(wb)
    sheet_wrong_headers(wb)
    sheet_title_then_table(wb)
    sheet_athlete_payroll(wb)
    return wb


def build_workbook_clean_vs_messy() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "规整列"
    ws.append(["姓名", "邮箱", "IP地址", "身份证", "银行卡号"])
    _style_header(ws)
    ws.append([NAMES[0], EMAILS[0], IPS[0], ID_CARDS[0], BANK_CARDS[0]])
    ws.append([NAMES[1], EMAILS[1], IPS[1], ID_CARDS[1], BANK_CARDS[1]])

    ws2 = wb.create_sheet("混乱散落")
    ws2.append(["业务线", "自由填写区", "其他"])
    _style_header(ws2)
    ws2.append(["网络", EMAILS[2], NAMES[2]])
    ws2.append(["财务", ID_CARDS[2], BANK_CARDS[2]])
    ws2.append(["人事", NAMES[3], IPS[2]])
    ws2.append(["行政", f"联系人{NAMES[5]}邮箱{EMAILS[3]}", "普通"])
    return wb


def build_workbook_in_cell_stress() -> Workbook:
    """格内扫描压力：多 PII、粘连、双邮箱、弱特征不应扫。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "格内压力"
    ws.append(["用例", "原文"])
    _style_header(ws)
    cases = [
        ("双邮箱", f"主送 {EMAILS[0]} 抄送 {EMAILS[4]}"),
        ("粘连邮箱", f"联系邮箱{EMAILS[1]}请回复"),
        ("四件套", f"{NAMES[0]} / {EMAILS[0]} / {IPS[0]} / {ID_CARDS[0]} / {BANK_CARDS[0]}"),
        ("英文句", f"Please contact {EMAILS[2]} from host {IPS[1]}."),
        ("手机号句中", f"紧急联系电话 {PHONES[0]}，邮箱 {EMAILS[3]}"),
        ("版本非IP", "客户端版本 v1.2.3 发布于 2024.1.1 金额 2.5"),
        ("公式文本", "公式保护 =SUM(H13:H17) 不改"),
        ("空", ""),
        ("仅人名句", f"今日值班{NAMES[2]}，替班{EXTRA_NAMES[0]}。"),
        ("长备注", "审计说明：" + "（无敏感）" * 8 + f" 联系人 {NAMES[5]} 邮箱 {EMAILS[0]} 完。"),
    ]
    for i, (name, text) in enumerate(cases, start=1):
        ws.append([name, text])
    ws.column_dimensions["B"].width = 80
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True)
    return wb


def build_workbook_false_positives() -> Workbook:
    """防误伤：列名过宽、日期小数、排除词、公式。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "列名过宽"
    ws.append(["备注邮箱", "联系", "金额", "版本", "类型"])
    _style_header(ws)
    ws.append([EMAILS[0], PHONES[0], "12000", "v2.3.1", "纪要"])
    ws.append(["无", "普通", "2.5", "2024.1.1", "工单"])
    ws.append(["策划部", "一队", "100", "1.0.0", "合计"])
    ws.append([f"见正文 {EMAILS[1]}", "普通", "0", "v1.2.3", "其他"])

    ws2 = wb.create_sheet("排除词整格")
    ws2.append(["字段A", "字段B"])
    _style_header(ws2)
    for w in ["合计", "策划部", "一队", "其他", "实习生", "备注"]:
        ws2.append([w, "普通文本"])
    return wb


def build_workbook_scatter() -> Workbook:
    """20 行 × 8 列，敏感值轮转落在不同列。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "散落矩阵"
    headers = ["列A", "列B", "列C", "列D", "列E", "列F", "列G", "列H"]
    ws.append(headers)
    _style_header(ws)
    payloads = [
        EMAILS[0],
        IPS[0],
        ID_CARDS[0],
        BANK_CARDS[0],
        NAMES[0],
        EXTRA_NAMES[0],
        f"备注：{EMAILS[1]}",
        "普通无敏感",
        PHONES[0],
        "2.5",
        NAMES[2],
        f"IP{IPS[3]}已封禁",
        ID_CARDS[3],
        f"{NAMES[1]}提交",
        BANK_CARDS[3],
        "策划部",
        EMAILS[4],
        "合计",
        f"证{ID_CARDS[1]}",
        EXTRA_NAMES[2],
    ]
    for i, payload in enumerate(payloads):
        row = ["—"] * 8
        row[i % 8] = payload
        ws.append(row)
    _widen(ws, 22)
    return wb


def build_workbook_audit_pack() -> Workbook:
    """模拟多 sheet 审计底稿。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "访谈纪要"
    ws.append(["序号", "纪要摘要"])
    _style_header(ws)
    ws.append([1, f"{NAMES[0]}表示系统账号为 {EMAILS[0]}，办公网 {IPS[0]}。"])
    ws.append([2, f"证人{NAMES[3]}出示身份证 {ID_CARDS[0]}。"])
    ws.append([3, "会议纪要无敏感字段。"])
    ws.append([4, f"备用联系 {EXTRA_NAMES[0]}（不在人员清单）。"])

    ws2 = wb.create_sheet("附件清单")
    ws2.append(["文件名", "备注"])
    _style_header(ws2)
    ws2.append(["scan001.pdf", f"含银行卡 {BANK_CARDS[0]}"])
    ws2.append(["mail.eml", f"发件人 {EMAILS[3]}"])
    ws2.append(["readme.txt", "无"])

    ws3 = wb.create_sheet("抽样表")
    ws3.append(["样本ID", "抽取说明", "结果"])
    _style_header(ws3)
    ws3.append(["S-01", f"邮箱列混填 {EMAILS[2]} 与 无", "待核"])
    ws3.append(["S-02", f"{NAMES[6]} / {PHONES[2]} / {IPS[2]}", "待核"])
    return wb


def write_scatter_csv(path: Path) -> None:
    lines = ["事项,自由填写"]
    lines += [
        f"通知,请联系 {EMAILS[0]}",
        f"网络,主机 {IPS[0]} 离线",
        f"入职,{NAMES[0]} 证件 {ID_CARDS[0]}",
        "金额,2.5",
        f"电话,手机 {PHONES[0]} 邮箱 {EMAILS[1]}",
        f"名单外,{EXTRA_NAMES[1]} 不在清单",
    ]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_readme(path: Path) -> None:
    path.write_text(
        """# 混乱布局测试数据（messy）

用来压当前表格引擎：列映射先校验、整格值级检测、格内强特征扫描、人员清单。

## 文件

| 文件 | 说明 |
|------|------|
| `messy_mixed.xlsx` | 混乱备注 / 嵌套正文 / 无语义列名 / 标题错位 / 选手收入 |
| `messy_vs_clean.xlsx` | 规整列 vs 混乱散落 |
| `in_cell_stress.xlsx` | 格内多 PII、粘连邮箱、手机号句中、版本/日期防误伤 |
| `false_positives.xlsx` | 「备注邮箱」列名过宽、排除词整格 |
| `scatter.xlsx` | 20 行敏感值轮转落在 A–H 列 |
| `audit_pack.xlsx` | 访谈纪要 / 附件清单 / 抽样表 |
| `scatter.csv` | CSV 对照 |
| `people.csv` | 清单内：张伟/李娜/欧阳修/司马光/东方不败/王芳/陈静 |

## 建议命令

```bash
maskit mask testdata/messy/in_cell_stress.xlsx --person-list testdata/messy/people.csv -o /tmp/out.xlsx
```

## 预期（有人员清单）

- 格内邮箱 / IP / 身份证 / 银行卡 → 遮
- 格内清单人名 → 遮；清单外赵磊/黄敏/周杰 → 不遮
- 句中手机号 → **不遮**（弱特征，防误伤）
- `无` / `策划部` 即使列名叫备注邮箱 → **不套** email 模板
- `2.5` / `2024.1.1` / `=SUM(...)` → 不遮
""",
        encoding="utf-8",
    )


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    write_people_csv(OUT / "people.csv")
    build_workbook_mixed().save(OUT / "messy_mixed.xlsx")
    build_workbook_clean_vs_messy().save(OUT / "messy_vs_clean.xlsx")
    build_workbook_in_cell_stress().save(OUT / "in_cell_stress.xlsx")
    build_workbook_false_positives().save(OUT / "false_positives.xlsx")
    build_workbook_scatter().save(OUT / "scatter.xlsx")
    build_workbook_audit_pack().save(OUT / "audit_pack.xlsx")
    write_scatter_csv(OUT / "scatter.csv")
    write_readme(OUT / "README.md")
    print(f"wrote {OUT}")
    for p in sorted(OUT.iterdir()):
        print(f"  {p.name} ({p.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
