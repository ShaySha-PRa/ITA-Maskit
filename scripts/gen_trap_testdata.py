"""对抗测试集：专门打当前引擎容易误伤/漏检的边界。

三套：
1. keyword_traps.xlsx  — 列名关键词过宽（地址/单位/版本/联系/编号/账号）
2. overlap_traps.xlsx  — 重叠 PII、子串人名、伪装 IP、JSON/HTML、空格分隔
3. layout_traps.xlsx   — 标题当表头、重复列名、超长格、数字型身份证

用法: python scripts/gen_trap_testdata.py
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, numbers

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "testdata" / "traps"

PEOPLE = ["张伟", "李娜", "欧阳修", "东方不败"]
EMAIL = "alice@corp.example"
IP = "10.1.2.3"
ID18 = "110101199003077777"
BANK19 = "6222021234567890123"


def _hdr(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.font = Font(bold=True)


def build_keyword_traps() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "列名陷阱"
    # 这些列名都会自动匹配到某条规则，但格子里经常不是那种值
    ws.append(["家庭地址", "单位", "版本", "联系人", "编号", "账号", "卡号"])
    _hdr(ws)
    ws.append([
        "上海市浦东新区",   # 地址→ip，值不是 IP
        "财务部",           # 单位→company，.+ 会整格遮
        "2024.1.1",         # 版本→app_version，日期像版本
        "张伟",             # 联系→phone，值是人名
        "2026001",          # 编号→employee_id，无连字符
        "admin",            # 账号→account
        BANK19,             # 卡号→bank_card，真卡号
    ])
    ws.append([
        IP,                 # 真 IP 落在「家庭地址」
        "亚玛芬体育",
        "v1.2.3",
        "赵磊",             # 清单外
        "EID-7001",         # 真工号形态
        "张伟",
        "不是卡号",
    ])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18
    return wb


def build_overlap_traps() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "重叠与伪装"
    ws.append(["用例", "原文"])
    _hdr(ws)
    rows = [
        ("子串人名", "选手张伟达获奖"),
        ("复姓词", "合唱歌曲东方红"),
        ("清单名在机构", "东方不败工作室来函"),
        ("双名粘连", "张伟李娜均出席"),
        ("版本伪装IP", "客户端 v10.20.30.40 发布"),
        ("JSON", '{"email":"alice@corp.example","host":"10.1.2.3"}'),
        ("HTML", "<p>联系 alice@corp.example</p>"),
        ("空格身份证", "身份证 110101 19900307 7777"),
        ("横线卡号", "卡号 6222-0212-3456-7890"),
        ("全角邮箱", "邮箱：ａｌｉｃｅ＠ｃｏｒｐ．ｅｘａｍｐｌｅ"),
        ("18位既像证又像卡", ID18),
        ("URL非邮箱", "见 https://corp.example/users/alice"),
        ("@不是邮箱", "请@全体成员 查看"),
        ("真邮箱无空格", f"联系人李娜邮箱{EMAIL}"),
        ("换行格", f"第一行\n邮箱 {EMAIL}\nIP {IP}"),
    ]
    for name, text in rows:
        ws.append([name, text])
    ws.column_dimensions["B"].width = 70
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True)
    return wb


def build_layout_traps() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "数字与精度"
    ws.append(["说明", "证件号", "备注"])
    _hdr(ws)
    # 文本型 18 位：应能遮
    ws.append(["文本身份证", ID18, "ok"])
    # Excel 数字精度只有 ~15 位，18 位会丢精度 → 几乎必然漏检
    r = ws.max_row + 1
    ws.append(["数字身份证", None, "excel 精度陷阱"])
    cell = ws.cell(r, 2, value=int(ID18))
    cell.number_format = numbers.FORMAT_NUMBER
    ws.append(["科学计数", 1.10101199003078e17, "同上"])

    ws2 = wb.create_sheet("重复列名")
    ws2.append(["备注", "备注", "说明"])
    _hdr(ws2)
    ws2.append([EMAIL, IP, "张伟"])

    ws3 = wb.create_sheet("超长夹带")
    ws3.append(["文档"])
    _hdr(ws3)
    ws3.append(["【审计底稿】" + ("无敏感。" * 80) + f" 最后才出现 {EMAIL} 与 {IP}。"])
    ws3.column_dimensions["A"].width = 40

    ws4 = wb.create_sheet("标题当表头")
    ws4.merge_cells("A1:C1")
    ws4["A1"] = "机密：抽查底稿"
    ws4.append([])
    ws4.append(["姓名", "邮箱", "IP"])
    ws4.append(["张伟", EMAIL, IP])
    return wb


def write_people(path: Path) -> None:
    path.write_text("name\n" + "\n".join(PEOPLE) + "\n", encoding="utf-8")


def write_readme(path: Path) -> None:
    path.write_text(
        """# 对抗测试集（traps）

故意设计成审计现场会翻车的形态，用来找误伤和漏检。

| 文件 | 陷阱 |
|------|------|
| `keyword_traps.xlsx` | 列名「地址/单位/版本/联系/编号/账号」绑错规则 |
| `overlap_traps.xlsx` | 张伟达含子串、东方红、v10.20.30.40、JSON/HTML、空格证号、全角邮箱 |
| `layout_traps.xlsx` | Excel 数字 18 位丢精度、重复列名、超长格末尾才出现 PII、标题行当表头 |
| `people.csv` | 张伟 李娜 欧阳修 东方不败（无赵磊） |

```bash
maskit mask testdata/traps/overlap_traps.xlsx --person-list testdata/traps/people.csv -o /tmp/trap.xlsx
```
""",
        encoding="utf-8",
    )


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    write_people(OUT / "people.csv")
    build_keyword_traps().save(OUT / "keyword_traps.xlsx")
    build_overlap_traps().save(OUT / "overlap_traps.xlsx")
    build_layout_traps().save(OUT / "layout_traps.xlsx")
    write_readme(OUT / "README.md")
    print(f"wrote {OUT}")
    for p in sorted(OUT.iterdir()):
        print(f"  {p.name} ({p.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
