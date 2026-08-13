"""第二轮复杂验证集：换人名/换场景，打上一轮修复后的组合边界。

产出 testdata/v2/
用法: python3 scripts/gen_v2_testdata.py
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, numbers

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "testdata" / "v2"

# 与 traps/messy 错开，避免「背答案」
PEOPLE = ["司马光", "上官婉儿", "诸葛孔明", "刘洋", "王芳", "Alice Chen"]
EMAIL = "nina.wu@audit.example"
EMAIL2 = "ops+oncall@audit.example"
IP = "172.16.9.41"
IP2 = "192.168.0.88"
ID18 = "44030119951212567X"
ID18_LOWER = "44030119951212567x"
BANK16 = "6222021234567890"
BANK_DASH = "6222-0212-3456-7890"
BANK_SPACE = "6222 0212 3456 7890"
FW_EMAIL = "ｎｉｎａ．ｗｕ＠ａｕｄｉｔ．ｅｘａｍｐｌｅ"


def _hdr(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.font = Font(bold=True)


def _widen(ws, width: int = 22) -> None:
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = width


def build_columns() -> Workbook:
    """列语义：修过的关键词 + 新列名。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "列语义"
    ws.append(["签约单位", "单位", "计量单位", "往来单位", "版本", "说明"])
    _hdr(ws)
    ws.append(["小菜园集团", "审计部", "元", "亚玛芬体育", "2023.12.31", "日期当版本"])
    ws.append(["MayAir", "人事处", "万元", "华为", "v2.0.1", "真版本"])
    ws.append(["Acme Inc", "法务科", "台", "腾讯", "1.4.2", "无v版本"])
    _widen(ws)
    return wb


def build_cocktail() -> Workbook:
    """一格混装 + 人名/IP/证卡新写法。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "一格混装"
    ws.append(["用例", "原文"])
    _hdr(ws)
    rows = [
        (
            "五件套句末句号",
            f"经办司马光；证 {ID18[:6]} {ID18[6:14]} {ID18[14:]}；"
            f"卡 {BANK_DASH}；邮 {FW_EMAIL}；主机 {IP}。",
        ),
        (
            "名单加邮箱加伪IP",
            f"请联系上官婉儿邮箱{EMAIL}，客户端 v8.8.8.8 勿改",
        ),
        ("二字加长", "选手刘洋达提交了申请"),
        ("机构后缀大学", "诸葛孔明大学来函"),
        ("三字加长", "司马光华不是清单名"),
        ("虚词的", "王芳的申请已受理"),
        ("双名粘连", "王芳刘洋均出席"),
        ("英文名", "On-call: Alice Chen approved."),
        ("清单外", "经办赵磊复核"),
        ("空格卡号", f"卡号 {BANK_SPACE} 已核"),
        ("小写X证", f"证件 {ID18_LOWER}"),
        ("加号邮箱", f"值班箱 {EMAIL2} 勿外发"),
        ("URL里的IP", f"后台 https://{IP2}/admin 请走堡垒"),
        ("端口IP", f"监听 {IP}:8080"),
        ("CIDR", f"网段 {IP2}/24 放行"),
        ("非法八位", "探测 999.1.1.1 失败"),
        ("广播地址", "网关 255.255.255.255 可达"),
        ("markdown邮箱", f"见 [联系](mailto:{EMAIL})"),
        ("JSON嵌套", '{"owner":"王芳","ip":"%s","mail":"%s"}' % (IP, EMAIL)),
        ("公式文本", "公式保护 =SUM(H13:H17) 不改"),
        ("真版本非IP", "客户端版本 v1.2.3 发布于 2024.1.1 金额 2.5"),
        ("手机不扫", "通知 13800001111 即可"),
    ]
    for name, text in rows:
        ws.append([name, text])
    ws.column_dimensions["B"].width = 80
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True)
    return wb


def build_layout() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "三重复列"
    ws.append(["说明", "说明", "说明"])
    _hdr(ws)
    ws.append([EMAIL, IP, "刘洋"])

    ws2 = wb.create_sheet("数字证号")
    ws2.append(["说明", "证件号"])
    _hdr(ws2)
    ws2.append(["文本X", ID18])
    r = ws2.max_row + 1
    ws2.append(["数字18位", None])
    cell = ws2.cell(r, 2, value=int("440301199512125678"))  # 末位用数字，模拟 Excel 存成 number
    cell.number_format = numbers.FORMAT_NUMBER

    ws3 = wb.create_sheet("标题当表头")
    ws3.merge_cells("A1:C1")
    ws3["A1"] = "抽查底稿（机密）"
    ws3.append([])
    ws3.append(["经办", "邮箱", "IP"])
    ws3.append(["王芳", EMAIL, IP])

    ws4 = wb.create_sheet("超长夹带")
    ws4.append(["文档"])
    _hdr(ws4)
    ws4.append(["【纪要】" + ("无敏感。" * 60) + f" 末尾 {EMAIL} / {IP} / 诸葛孔明。"])
    return wb


def write_csv(path: Path) -> None:
    lines = [
        "事项,备注",
        f"巡检,主机 {IP} 联系 {EMAIL}",
        "部门,审计部",
        f"入职,刘洋 证件 {ID18}",
        "版本,v8.8.8.8 发布说明",
        f"对账,卡 {BANK_DASH}",
    ]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_people(path: Path) -> None:
    path.write_text("name\n" + "\n".join(PEOPLE) + "\n", encoding="utf-8")


def write_readme(path: Path) -> None:
    path.write_text(
        """# 第二轮复杂验证集（v2）

与 testdata/traps、testdata/messy 错开人名和场景，用来复验上一轮误伤/漏检修复。

人员清单：司马光 上官婉儿 诸葛孔明 刘洋 王芳 Alice Chen（无赵磊）

```bash
python3 scripts/gen_v2_testdata.py
python3 -m pytest tests/test_v2_testdata.py -v
```
""",
        encoding="utf-8",
    )


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    write_people(OUT / "people.csv")
    build_columns().save(OUT / "columns.xlsx")
    build_cocktail().save(OUT / "cocktail.xlsx")
    build_layout().save(OUT / "layout.xlsx")
    write_csv(OUT / "scatter.csv")
    write_readme(OUT / "README.md")
    print(f"wrote {OUT}")
    for p in sorted(OUT.iterdir()):
        print(f"  {p.name} ({p.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
