"""用混乱布局测试集回归当前表格引擎（列校验 + 格内扫描 + 人员清单）。"""
from __future__ import annotations

import importlib.util
from pathlib import Path

import pytest
from openpyxl import load_workbook

from maskit.io import mask_file
from maskit.rules.loader import load_ruleset
from maskit.rules.name_company import load_person_list

_GEN = Path(__file__).resolve().parents[1] / "scripts" / "gen_messy_testdata.py"
_spec = importlib.util.spec_from_file_location("gen_messy_testdata", _GEN)
assert _spec and _spec.loader
gen = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(gen)


@pytest.fixture
def people() -> set[str]:
    return set(gen.NAMES)


def _mask_wb(wb_builder, tmp_path, people, name: str):
    src = tmp_path / f"{name}.xlsx"
    out = tmp_path / f"{name}_masked.xlsx"
    wb_builder().save(src)
    mask_file(str(src), str(out), load_ruleset(), None, person_list=people)
    return load_workbook(out, data_only=True)


def _col(ws, header: str) -> int:
    for i, cell in enumerate(ws[1], start=1):
        if cell.value == header:
            return i
    raise AssertionError(f"no column {header}")


def test_in_cell_stress_with_person_list(tmp_path, people):
    wb = _mask_wb(gen.build_workbook_in_cell_stress, tmp_path, people, "stress")
    ws = wb["格内压力"]
    by_case = {row[0]: row[1] for row in ws.iter_rows(min_row=2, values_only=True)}

    assert gen.EMAILS[0] not in by_case["双邮箱"]
    assert gen.EMAILS[4] not in by_case["双邮箱"]
    assert "a***@corp.example" in by_case["双邮箱"]

    assert gen.EMAILS[1] not in by_case["粘连邮箱"]
    assert "b***@company.cn" in by_case["粘连邮箱"]

    four = by_case["四件套"]
    assert gen.EMAILS[0] not in four
    assert gen.IPS[0] not in four
    assert gen.ID_CARDS[0] not in four
    assert gen.BANK_CARDS[0] not in four
    assert "张*" in four

    assert gen.EMAILS[2] not in by_case["英文句"]
    assert gen.IPS[1] not in by_case["英文句"]

    phone_row = by_case["手机号句中"]
    assert gen.PHONES[0] in phone_row  # 弱特征格内不扫
    assert gen.EMAILS[3] not in phone_row

    assert by_case["版本非IP"] == "客户端版本 v1.2.3 发布于 2024.1.1 金额 2.5"
    assert by_case["公式文本"] == "公式保护 =SUM(H13:H17) 不改"

    only_names = by_case["仅人名句"]
    assert "欧阳修" not in only_names
    assert "赵磊" in only_names  # 清单外

    long = by_case["长备注"]
    assert gen.EMAILS[0] not in long
    assert "王芳" not in long


def test_false_positive_email_column(tmp_path, people):
    wb = _mask_wb(gen.build_workbook_false_positives, tmp_path, people, "fp")
    ws = wb["列名过宽"]
    emails = [row[0] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert gen.EMAILS[0] not in emails[0]
    assert emails[1] == "无"
    assert emails[2] == "策划部"
    assert gen.EMAILS[1] not in str(emails[3])

    types = [row[4] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert types[0] == "纪要"  # 有清单关闭启发式
    assert types[2] == "合计"

    ws2 = wb["排除词整格"]
    kept = [row[0] for row in ws2.iter_rows(min_row=2, values_only=True)]
    assert kept == ["合计", "策划部", "一队", "其他", "实习生", "备注"]


def test_scatter_rotating_columns(tmp_path, people):
    wb = _mask_wb(gen.build_workbook_scatter, tmp_path, people, "scatter")
    ws = wb["散落矩阵"]
    blob = " ".join(
        str(c) if c is not None else ""
        for row in ws.iter_rows(min_row=2, values_only=True)
        for c in row
    )
    assert gen.EMAILS[0] not in blob
    assert gen.EMAILS[1] not in blob
    assert gen.EMAILS[4] not in blob
    assert gen.IPS[0] not in blob
    assert gen.ID_CARDS[0] not in blob
    assert gen.BANK_CARDS[0] not in blob
    assert "张伟" not in blob
    assert "赵磊" in blob
    assert "周杰" in blob
    assert gen.PHONES[0] in blob  # 整格手机号：phone 不在值级白名单
    assert "2.5" in blob
    assert "策划部" in blob


def test_audit_pack_minutes(tmp_path, people):
    wb = _mask_wb(gen.build_workbook_audit_pack, tmp_path, people, "audit")
    notes = [row[1] for row in wb["访谈纪要"].iter_rows(min_row=2, values_only=True)]
    assert gen.EMAILS[0] not in notes[0]
    assert gen.IPS[0] not in notes[0]
    assert "张伟" not in notes[0]
    assert gen.ID_CARDS[0] not in notes[1]
    assert "司马光" not in notes[1]
    assert notes[2] == "会议纪要无敏感字段。"
    assert "赵磊" in notes[3]

    att = [row[1] for row in wb["附件清单"].iter_rows(min_row=2, values_only=True)]
    assert gen.BANK_CARDS[0] not in att[0]
    assert gen.EMAILS[3] not in att[1]
    assert att[2] == "无"


def test_mixed_nested_and_payroll(tmp_path, people):
    wb = _mask_wb(gen.build_workbook_mixed, tmp_path, people, "mixed")
    prose = [row[2] for row in wb["嵌套正文"].iter_rows(min_row=2, values_only=True)]
    assert gen.EMAILS[0] not in prose[0]
    assert "张伟" not in prose[0]
    assert gen.IPS[0] not in prose[1]
    assert "李娜" not in prose[1]
    assert prose[4] == "本周无敏感事项，仅例会纪要。"

    pay = list(wb["选手收入"].iter_rows(min_row=2, values_only=True))
    assert pay[2][3] == "赵磊"
    assert pay[2][5] == "无"
    assert pay[3][5] == "策划部"
    assert gen.EMAILS[0] not in str(pay[0][5])


def test_scatter_csv(tmp_path, people):
    src = tmp_path / "scatter.csv"
    out = tmp_path / "scatter_masked.csv"
    gen.write_scatter_csv(src)
    mask_file(str(src), str(out), load_ruleset(), None, person_list=people)
    text = out.read_text(encoding="utf-8")
    assert gen.EMAILS[0] not in text
    assert gen.IPS[0] not in text
    assert gen.ID_CARDS[0] not in text
    assert "张伟" not in text
    assert gen.PHONES[0] in text
    assert "黄敏" in text
    assert "2.5" in text
