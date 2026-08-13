"""对抗测试：审计上应有的对错（误伤/漏检修复后的硬断言）。"""
from __future__ import annotations

import importlib.util
from pathlib import Path

from openpyxl import load_workbook

from maskit.io import mask_file
from maskit.rules.loader import load_ruleset

_GEN = Path(__file__).resolve().parents[1] / "scripts" / "gen_trap_testdata.py"
_spec = importlib.util.spec_from_file_location("gen_trap_testdata", _GEN)
assert _spec and _spec.loader
gen = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(gen)

PEOPLE = set(gen.PEOPLE)


def _mask(builder, tmp_path, name: str):
    src = tmp_path / f"{name}.xlsx"
    out = tmp_path / f"{name}_out.xlsx"
    builder().save(src)
    mask_file(str(src), str(out), load_ruleset(), None, person_list=PEOPLE)
    return load_workbook(out, data_only=True)


def test_keyword_traps_what_engine_gets_right(tmp_path):
    ws = _mask(gen.build_keyword_traps, tmp_path, "kw")["列名陷阱"]
    headers = [c.value for c in ws[1]]
    r1 = dict(zip(headers, [c.value for c in ws[2]]))
    r2 = dict(zip(headers, [c.value for c in ws[3]]))

    assert r1["家庭地址"] == "上海市浦东新区"
    assert r1["编号"] == "2026001"
    assert "张*" in str(r1["联系人"])
    assert gen.BANK19 not in str(r1["卡号"])
    assert str(r1["账号"]) != "admin"
    assert "*.*.*.*" in str(r2["家庭地址"])
    assert r2["联系人"] == "赵磊"
    assert r2["卡号"] == "不是卡号"
    assert str(r2["编号"]) != "EID-7001"


def test_trap_department_not_company(tmp_path):
    ws = _mask(gen.build_keyword_traps, tmp_path, "kw")["列名陷阱"]
    r1 = dict(zip([c.value for c in ws[1]], [c.value for c in ws[2]]))
    assert r1["单位"] == "财务部"


def test_trap_date_not_app_version(tmp_path):
    ws = _mask(gen.build_keyword_traps, tmp_path, "kw")["列名陷阱"]
    r1 = dict(zip([c.value for c in ws[1]], [c.value for c in ws[2]]))
    assert r1["版本"] == "2024.1.1"


def test_overlap_hits_json_html_newline(tmp_path):
    ws = _mask(gen.build_overlap_traps, tmp_path, "ov")["重叠与伪装"]
    c = {r[0]: (r[1] or "") for r in ws.iter_rows(min_row=2, values_only=True)}
    assert "alice@corp.example" not in c["JSON"]
    assert "10.1.2.3" not in c["JSON"]
    assert "alice@corp.example" not in c["HTML"]
    assert "alice@corp.example" not in c["换行格"]
    assert "10.1.2.3" not in c["换行格"]
    assert "alice@corp.example" not in c["真邮箱无空格"]
    assert "李娜" not in c["真邮箱无空格"]
    assert "东方红" in c["复姓词"]
    assert "张*" in c["双名粘连"] and "李*" in c["双名粘连"]
    assert "corp.example/users" in c["URL非邮箱"]
    assert c["@不是邮箱"] == "请@全体成员 查看"
    assert gen.ID18 not in c["18位既像证又像卡"]


def test_trap_name_substring(tmp_path):
    ws = _mask(gen.build_overlap_traps, tmp_path, "ov")["重叠与伪装"]
    c = {r[0]: (r[1] or "") for r in ws.iter_rows(min_row=2, values_only=True)}
    assert "张伟达" in c["子串人名"]


def test_trap_name_inside_org(tmp_path):
    ws = _mask(gen.build_overlap_traps, tmp_path, "ov")["重叠与伪装"]
    c = {r[0]: (r[1] or "") for r in ws.iter_rows(min_row=2, values_only=True)}
    assert "东方不败工作室" in c["清单名在机构"]


def test_trap_version_looks_like_ip(tmp_path):
    ws = _mask(gen.build_overlap_traps, tmp_path, "ov")["重叠与伪装"]
    c = {r[0]: (r[1] or "") for r in ws.iter_rows(min_row=2, values_only=True)}
    assert "v10.20.30.40" in c["版本伪装IP"]


def test_trap_spaced_id_card(tmp_path):
    ws = _mask(gen.build_overlap_traps, tmp_path, "ov")["重叠与伪装"]
    c = {r[0]: (r[1] or "") for r in ws.iter_rows(min_row=2, values_only=True)}
    assert "110101199003077777" not in c["空格身份证"].replace(" ", "")


def test_trap_dashed_bank_card(tmp_path):
    ws = _mask(gen.build_overlap_traps, tmp_path, "ov")["重叠与伪装"]
    c = {r[0]: (r[1] or "") for r in ws.iter_rows(min_row=2, values_only=True)}
    digits = "".join(ch for ch in c["横线卡号"] if ch.isdigit())
    assert "6222021234567890" not in digits


def test_trap_fullwidth_email(tmp_path):
    ws = _mask(gen.build_overlap_traps, tmp_path, "ov")["重叠与伪装"]
    c = {r[0]: (r[1] or "") for r in ws.iter_rows(min_row=2, values_only=True)}
    assert "ａｌｉｃｅ" not in c["全角邮箱"]


def test_layout_text_id_and_long_cell(tmp_path):
    src = tmp_path / "ly.xlsx"
    out = tmp_path / "ly_out.xlsx"
    wb = gen.build_layout_traps()
    wb.save(src)
    mask_file(str(src), str(out), load_ruleset(), None, person_list=PEOPLE)
    out_wb = load_workbook(out, data_only=True)

    rows = list(out_wb["数字与精度"].iter_rows(min_row=2, values_only=True))
    assert gen.ID18 not in str(rows[0][1])

    long = next(out_wb["超长夹带"].iter_rows(min_row=2, values_only=True))[0]
    assert gen.EMAIL not in str(long)
    assert gen.IP not in str(long)

    blob = " ".join(
        str(c) for row in out_wb["标题当表头"].iter_rows(values_only=True) for c in row if c
    )
    assert gen.EMAIL not in blob
    assert gen.IP not in blob


def test_trap_excel_numeric_id(tmp_path):
    """数字型 18 位：禁止科学计数；若能读成 18 位则必须已遮。精度损失无法还原原文。"""
    src = tmp_path / "ly.xlsx"
    out = tmp_path / "ly_out.xlsx"
    wb = gen.build_layout_traps()
    wb.save(src)
    mask_file(str(src), str(out), load_ruleset(), None, person_list=PEOPLE)
    rows = list(load_workbook(out, data_only=True)["数字与精度"].iter_rows(min_row=2, values_only=True))
    num_id = str(rows[1][1] or "")
    assert "e+" not in num_id.lower()
    digits = "".join(ch for ch in num_id if ch.isdigit())
    if len(digits) == 18:
        assert gen.ID18 not in num_id
        assert "*" in num_id


def test_trap_duplicate_headers_renamed(tmp_path):
    """重复列名去重为 备注 / 备注_2，两列 PII 都能遮。"""
    src = tmp_path / "dup.xlsx"
    out = tmp_path / "dup_out.xlsx"
    wb = gen.build_layout_traps()
    for name in list(wb.sheetnames):
        if name != "重复列名":
            del wb[name]
    wb.save(src)
    mask_file(str(src), str(out), load_ruleset(), None, person_list=PEOPLE)
    ws = load_workbook(out, data_only=True)["重复列名"]
    headers = [c.value for c in ws[1]]
    assert headers[0] == "备注"
    assert headers[1] == "备注_2"
    row = [c.value for c in ws[2]]
    assert gen.EMAIL not in str(row[0])
    assert gen.IP not in str(row[1])
    assert "张*" in str(row[2])
