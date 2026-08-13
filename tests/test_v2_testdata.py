"""第二轮复杂验证集：换人名/换场景，复验误伤与漏检修复。"""
from __future__ import annotations

import importlib.util
from pathlib import Path

import polars as pl
import pytest
from openpyxl import load_workbook

from maskit.io import mask_file
from maskit.rules.engine import apply_rules
from maskit.rules.loader import load_ruleset

_GEN = Path(__file__).resolve().parents[1] / "scripts" / "gen_v2_testdata.py"
_spec = importlib.util.spec_from_file_location("gen_v2_testdata", _GEN)
assert _spec and _spec.loader
gen = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(gen)

PEOPLE = set(gen.PEOPLE)


def _mask(builder, tmp_path, name: str):
    src = tmp_path / f"{name}.xlsx"
    out = tmp_path / f"{name}_out.xlsx"
    builder().save(src)
    mask_file(str(src), str(out), load_ruleset(), None, person_list=PEOPLE)
    return load_workbook(out, data_only=True)


def _cases(ws):
    return {r[0]: (r[1] or "") for r in ws.iter_rows(min_row=2, values_only=True)}


def test_v2_column_semantics(tmp_path):
    ws = _mask(gen.build_columns, tmp_path, "col")["列语义"]
    headers = [c.value for c in ws[1]]
    r1 = dict(zip(headers, [c.value for c in ws[2]]))
    r2 = dict(zip(headers, [c.value for c in ws[3]]))
    r3 = dict(zip(headers, [c.value for c in ws[4]]))

    assert r1["单位"] == "审计部"
    assert r1["计量单位"] == "元"
    assert r1["版本"] == "2023.12.31"
    assert r1["签约单位"] != "小菜园集团"
    assert str(r1["签约单位"]).startswith("小")
    assert r1["往来单位"] != "亚玛芬体育"

    assert r2["单位"] == "人事处"
    assert r2["版本"] == "v2.*.*"
    assert r2["往来单位"] != "华为"

    assert r3["单位"] == "法务科"
    assert r3["版本"] == "v1.*.*"


def test_v2_cocktail_hits_and_guards(tmp_path):
    c = _cases(_mask(gen.build_cocktail, tmp_path, "ck")["一格混装"])

    five = c["五件套句末句号"]
    assert "司马光" not in five
    assert "44030119951212567X" not in five.replace(" ", "")
    assert "6222021234567890" not in "".join(ch for ch in five if ch.isdigit())
    assert "ｎｉｎａ" not in five
    assert gen.IP not in five
    assert "*.*.*.*" in five

    combo = c["名单加邮箱加伪IP"]
    assert "上官婉儿" not in combo
    assert gen.EMAIL not in combo
    assert "v8.8.8.8" in combo

    assert "刘洋达" in c["二字加长"]
    assert "诸葛孔明大学" in c["机构后缀大学"]
    assert "王芳的申请" not in c["虚词的"] or "王*" in c["虚词的"]
    assert "王芳" not in c["虚词的"]
    assert "王*" in c["双名粘连"] and "刘*" in c["双名粘连"]
    assert "Alice Chen" not in c["英文名"]
    assert "赵磊" in c["清单外"]

    digits = "".join(ch for ch in c["空格卡号"] if ch.isdigit())
    assert "6222021234567890" not in digits
    assert "44030119951212567" not in c["小写X证"].lower().replace(" ", "")
    assert gen.EMAIL2 not in c["加号邮箱"]
    assert gen.IP2 not in c["URL里的IP"]
    assert gen.IP not in c["端口IP"]
    assert gen.IP2 not in c["CIDR"]
    assert "999.1.1.1" in c["非法八位"]
    assert "255.255.255.255" not in c["广播地址"]
    assert gen.EMAIL not in c["markdown邮箱"]
    assert gen.EMAIL not in c["JSON嵌套"]
    assert gen.IP not in c["JSON嵌套"]
    assert "王芳" not in c["JSON嵌套"]
    assert c["公式文本"] == "公式保护 =SUM(H13:H17) 不改"
    assert c["真版本非IP"] == "客户端版本 v1.2.3 发布于 2024.1.1 金额 2.5"
    assert "13800001111" in c["手机不扫"]


def test_v2_formula_skipped_in_dataframe():
    """整格以 = 开头的公式不扫描（Excel 真公式读出来是空，这里用 DataFrame 测引擎）。"""
    from maskit.rules.defs import RuleSet

    rs = RuleSet(defs=load_ruleset().defs, specs=[])
    df = pl.DataFrame({"备注": ["=SUM(H13:H17)", f"联系 {gen.EMAIL}"]})
    masked, count = apply_rules(df, rs, None, person_list=PEOPLE)
    rows = masked["备注"].to_list()
    assert rows[0] == "=SUM(H13:H17)"
    assert gen.EMAIL not in rows[1]
    assert count == 1


@pytest.mark.xfail(reason="更长姓名否决只覆盖 2 字清单名，三字「司马光」仍会切「司马光华」")
def test_v2_three_char_name_not_cut_longer(tmp_path):
    """三字清单名不应切四字：司马光华。"""
    c = _cases(_mask(gen.build_cocktail, tmp_path, "ck")["一格混装"])
    assert "司马光华" in c["三字加长"]


def test_v2_layout_and_csv(tmp_path):
    src = tmp_path / "ly.xlsx"
    out = tmp_path / "ly_out.xlsx"
    gen.build_layout().save(src)
    mask_file(str(src), str(out), load_ruleset(), None, person_list=PEOPLE)
    wb = load_workbook(out, data_only=True)

    dup = wb["三重复列"]
    headers = [c.value for c in dup[1]]
    assert headers[0] == "说明"
    assert headers[1] == "说明_2"
    assert headers[2] == "说明_3"
    row = [c.value for c in dup[2]]
    assert gen.EMAIL not in str(row[0])
    assert gen.IP not in str(row[1])
    assert "刘*" in str(row[2])

    ids = list(wb["数字证号"].iter_rows(min_row=2, values_only=True))
    assert gen.ID18 not in str(ids[0][1]).upper()
    num_id = str(ids[1][1] or "")
    assert "e+" not in num_id.lower()

    blob = " ".join(
        str(c) for row in wb["标题当表头"].iter_rows(values_only=True) for c in row if c
    )
    assert gen.EMAIL not in blob
    assert gen.IP not in blob

    long = next(wb["超长夹带"].iter_rows(min_row=2, values_only=True))[0]
    assert gen.EMAIL not in str(long)
    assert gen.IP not in str(long)
    assert "诸葛孔明" not in str(long)

    csv_src = tmp_path / "scatter.csv"
    csv_out = tmp_path / "scatter_out.csv"
    gen.write_csv(csv_src)
    mask_file(str(csv_src), str(csv_out), load_ruleset(), None, person_list=PEOPLE)
    text = csv_out.read_text(encoding="utf-8")
    assert gen.EMAIL not in text
    assert gen.IP not in text
    assert gen.ID18 not in text
    assert "6222021234567890" not in text.replace("-", "")
    assert "审计部" in text
    assert "v8.8.8.8" in text
